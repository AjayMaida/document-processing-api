import io
from pathlib import Path
from unittest.mock import MagicMock

import pymupdf
import pytest
from docx import Document as DocxDocument
from fastapi import HTTPException

from app.core.config import settings
from app.core.storage.local import LocalStorage
from app.models.document import Document
from app.services.exceptions import TextExtractionError
from app.services.text_extraction_service import TextExtractionService


def create_service(storage=None) -> tuple[TextExtractionService, MagicMock]:
    repository = MagicMock()
    if storage is None:
        storage = LocalStorage()
    service = TextExtractionService(repository=repository, storage=storage)
    return service, repository


def test_extract_txt_file(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    upload_dir = tmp_path / "uploads"
    upload_dir.mkdir()

    monkeypatch.setattr(settings, "upload_dir", str(upload_dir))
    storage = LocalStorage()

    document = Document(
        id=1,
        original_filename="example.txt",
        stored_filename="stored-example.txt",
        status="uploaded",
        user_id=1,
    )

    source_text = "Hello, this is extracted text."
    storage.upload(io.BytesIO(source_text.encode("utf-8")), document.stored_filename)

    service, repository = create_service(storage=storage)

    extracted_record = MagicMock()
    extracted_record.document_id = document.id
    repository.get_by_document_id.return_value = extracted_record

    result = service.extract_text(document)

    expected_path = f"extracted/{document.id}/extracted.txt"
    assert storage.exists(expected_path)
    output_stream = storage.download(expected_path)
    assert output_stream.read().decode("utf-8") == source_text

    assert result is extracted_record
    repository.update_status.assert_any_call(extracted_record, "processing")
    repository.update_status.assert_any_call(extracted_record, "completed")
    repository.update_text_path.assert_called_once_with(extracted_record, expected_path)
    assert repository.commit.call_count == 2


def test_extract_txt_file_missing_source(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    upload_dir = tmp_path / "uploads"
    upload_dir.mkdir()

    monkeypatch.setattr(settings, "upload_dir", str(upload_dir))
    storage = LocalStorage()

    document = Document(
        id=1,
        original_filename="missing.txt",
        stored_filename="missing.txt",
        status="uploaded",
        user_id=1,
    )

    service, repository = create_service(storage=storage)

    with pytest.raises(HTTPException):
        service.extract_text(document)

    repository.get_by_document_id.assert_not_called()


def test_extract_unsupported_file_type() -> None:
    service, _ = create_service()
    stream = io.BytesIO(b"name,value\nfoo,bar")

    with pytest.raises(HTTPException):
        service._extract_from_stream(stream, "example.csv")


def test_extract_txt_encoding_error() -> None:
    service, _ = create_service()
    # Invalid UTF-8 bytes in stream
    stream = io.BytesIO(b"\xff\xfe\xfd")

    with pytest.raises(TextExtractionError):
        service._extract_from_stream(stream, "invalid.txt")


def test_extract_pdf_file(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    upload_dir = tmp_path / "uploads"
    upload_dir.mkdir()

    monkeypatch.setattr(settings, "upload_dir", str(upload_dir))
    storage = LocalStorage()

    document = Document(
        id=2,
        original_filename="example.pdf",
        stored_filename="stored-example.pdf",
        status="uploaded",
        user_id=1,
    )

    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Hello from a PDF document.")
    pdf_bytes = pdf.tobytes()
    pdf.close()

    storage.upload(io.BytesIO(pdf_bytes), document.stored_filename)

    service, repository = create_service(storage=storage)

    extracted_record = MagicMock()
    extracted_record.document_id = document.id
    repository.get_by_document_id.return_value = extracted_record

    result = service.extract_text(document)

    expected_path = f"extracted/{document.id}/extracted.txt"
    assert storage.exists(expected_path)
    output_stream = storage.download(expected_path)
    extracted_content = output_stream.read().decode("utf-8")
    assert "Hello from a PDF document." in extracted_content

    assert result is extracted_record
    repository.update_status.assert_any_call(extracted_record, "processing")
    repository.update_status.assert_any_call(extracted_record, "completed")
    repository.update_text_path.assert_called_once_with(extracted_record, expected_path)


def test_extract_docx_file(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    upload_dir = tmp_path / "uploads"
    upload_dir.mkdir()

    monkeypatch.setattr(settings, "upload_dir", str(upload_dir))
    storage = LocalStorage()

    document = Document(
        id=3,
        original_filename="example.docx",
        stored_filename="stored-example.docx",
        status="uploaded",
        user_id=1,
    )

    docx = DocxDocument()
    docx.add_paragraph("Hello from a DOCX document.")
    docx.add_paragraph("This is the second paragraph.")
    docx_io = io.BytesIO()
    docx.save(docx_io)
    docx_io.seek(0)

    storage.upload(docx_io, document.stored_filename)

    service, repository = create_service(storage=storage)

    extracted_record = MagicMock()
    extracted_record.document_id = document.id
    repository.get_by_document_id.return_value = extracted_record

    result = service.extract_text(document)

    expected_path = f"extracted/{document.id}/extracted.txt"
    assert storage.exists(expected_path)
    output_stream = storage.download(expected_path)
    extracted_content = output_stream.read().decode("utf-8")

    assert "Hello from a DOCX document." in extracted_content
    assert "This is the second paragraph." in extracted_content

    assert result is extracted_record
    repository.update_status.assert_any_call(extracted_record, "processing")
    repository.update_status.assert_any_call(extracted_record, "completed")
    repository.update_text_path.assert_called_once_with(extracted_record, expected_path)
